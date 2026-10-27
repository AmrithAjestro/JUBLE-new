import re
from io import BytesIO
from pathlib import Path
from textwrap import wrap
from xml.sax.saxutils import escape as xml_escape

import streamlit as st
from docx import Document
from PyPDF2 import PdfReader

from dataset_loader import article_refs, coi_lookup
from module.llm import CHAT_MODEL, chat, extract_json

PROJECT_ROOT = Path(__file__).resolve().parent.parent
COMPLAINTS_FILE = PROJECT_ROOT / "complaints.txt"

# ------------------------------
# Legal relevance detection
# ------------------------------
_LAW_TERMS = re.compile(
    r"\b(ipc|bns|bnss|crpc|indian penal code|constitution|fundamental (right|duty)|article|acts?|legal|laws?|"
    r"court|justice|section|motor vehicles act|mv act|traffic|drunk(en)? driving|drink and drive|dui|"
    r"fines?|penalty|challan|wrong side driving|over ?speeding|wrong parking|bail|fir|arrest|police|"
    r"petition|writ|lawyer|advocate|judge|offen[cs]e|crime|criminal|civil|contract|property|divorce)\b",
    re.I,
)
_FINE_TERMS = re.compile(r"\b(fines?|penalty|penalties|challan|punishment|punishable)\b", re.I)


def is_law_related(text: str) -> bool:
    """Detects legal context (IPC/BNS, Constitution, common statutes, traffic fines)."""
    return bool(_LAW_TERMS.search(text or ""))


# ------------------------------
# PDF processing & summarization
# ------------------------------
def read_pdf(file) -> str:
    """Extracts all text from a PDF ('' on failure)."""
    try:
        return "\n\n".join((page.extract_text() or "").strip() for page in PdfReader(file).pages).strip()
    except Exception as e:
        st.error(f"❌ Error reading PDF: {e}")
        return ""


def generate_summary(long_text: str, chunk_chars: int = 8000) -> str:
    """Summarizes long legal text chunk by chunk, merging only when there is more than one chunk."""
    try:
        summaries = [
            chat([
                {"role": "system", "content": "You are a senior Indian legal expert. Summarize this content clearly."},
                {"role": "user", "content": chunk},
            ], temperature=0.2, max_tokens=1200)
            for chunk in wrap(long_text, chunk_chars)
        ]
        if len(summaries) == 1:
            return summaries[0]
        return chat([
            {"role": "system", "content": "Combine all partial summaries into one well-structured summary."},
            {"role": "user", "content": "\n\n".join(summaries)},
        ], temperature=0.2, max_tokens=1500)
    except Exception as e:
        return str(e)


# ------------------------------
# Export chat (ReportLab PDF with light Markdown rendering / DOCX)
# ------------------------------
def _find_font(names: list[str]) -> str | None:
    for path in names:
        if Path(path).exists():
            return path
    return None


def export_chat(format_type="pdf"):
    """Exports chat as PDF or DOCX. PDF mimics chat formatting (headings, bold/italics, lists, code blocks)."""
    messages = st.session_state.get("law_messages", [])
    if not messages:
        st.warning("⚠️ No chat messages to export.")
        return

    if format_type == "docx":
        buf = BytesIO()
        doc = Document()
        doc.add_heading("Law Chat Conversation", level=1)
        for msg in messages:
            doc.add_paragraph(f"{msg['role'].capitalize()}: {msg['content']}")
        doc.save(buf)
        st.download_button(
            "📘 Download as DOCX",
            data=buf.getvalue(),
            file_name="law_chat.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        return

    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, Preformatted, SimpleDocTemplate, Spacer

    # --- Fonts (Unicode + mono if available)
    fonts_dir = PROJECT_ROOT / "assets" / "fonts"
    uni_font, mono_font = "Helvetica", "Courier"
    try:
        uni_path = _find_font([
            fonts_dir / "DejaVuSans.ttf", r"C:\Windows\Fonts\DejaVuSans.ttf", r"C:\Windows\Fonts\Nirmala.ttf",
            r"C:\Windows\Fonts\ARIALUNI.TTF", r"C:\Windows\Fonts\SegoeUI.ttf",
        ])
        if uni_path:
            pdfmetrics.registerFont(TTFont("Uni", str(uni_path)))
            uni_font = "Uni"
        mono_path = _find_font([
            fonts_dir / "DejaVuSansMono.ttf", r"C:\Windows\Fonts\DejaVuSansMono.ttf", r"C:\Windows\Fonts\consola.ttf",
        ])
        if mono_path:
            pdfmetrics.registerFont(TTFont("UniMono", str(mono_path)))
            mono_font = "UniMono"
    except Exception:
        pass

    # --- Styles
    para = ParagraphStyle("para", parent=getSampleStyleSheet()["Normal"], fontName=uni_font, fontSize=11,
                          leading=15, alignment=TA_LEFT, spaceBefore=2, spaceAfter=4, wordWrap="CJK")
    h1 = ParagraphStyle("h1", parent=para, fontSize=15, leading=19, spaceBefore=8, spaceAfter=6)
    h2 = ParagraphStyle("h2", parent=para, fontSize=13, leading=17, spaceBefore=6, spaceAfter=4)
    title_style = ParagraphStyle("title", parent=para, fontSize=16, leading=20, spaceAfter=8)
    role_style = ParagraphStyle("role", parent=para, backColor=colors.whitesmoke, spaceBefore=4,
                                spaceAfter=2, fontSize=10.5)
    code_style = ParagraphStyle("code", parent=para, fontName=mono_font, backColor=colors.HexColor("#F6F8FA"),
                                fontSize=10.5, leading=14, leftIndent=4, rightIndent=4, spaceBefore=4, spaceAfter=6)

    # --- Tiny Markdown → Flowables
    code_fence = re.compile(r"```[^\n]*\n(.*?)```", re.DOTALL)
    inline_code = re.compile(r"`([^`]+)`")
    bold = re.compile(r"\*\*([^\*]+)\*\*")
    ital = re.compile(r"(?<!\*)\*([^\*]+)\*(?!\*)")
    list_item = re.compile(r"^\s*([-*]|\d+\.)\s+")

    def inline_markup(s: str) -> str:
        s = bold.sub(r"<b>\1</b>", xml_escape(s))
        s = ital.sub(r"<i>\1</i>", s)
        return inline_code.sub(lambda m: f'<font name="{mono_font}">{m.group(1)}</font>', s)

    def paragraphs_and_lists(block: str):
        flow, buf = [], []

        def flush():
            text = "\n".join(buf).strip()
            buf.clear()
            if not text:
                return
            if text.startswith("## "):
                flow.append(Paragraph(inline_markup(text[3:]), h2))
            elif text.startswith("# "):
                flow.append(Paragraph(inline_markup(text[2:]), h1))
            else:
                flow.append(Paragraph(inline_markup(text), para))

        lines = block.splitlines()
        i = 0
        while i < len(lines):
            if list_item.match(lines[i]):
                flush()
                ordered = bool(re.match(r"^\s*\d+\.", lines[i]))
                items = []
                while i < len(lines) and list_item.match(lines[i]):
                    items.append(ListItem(Paragraph(inline_markup(list_item.sub("", lines[i]).rstrip()), para)))
                    i += 1
                flow.append(ListFlowable(items, bulletType="1" if ordered else "bullet", start="1",
                                         leftIndent=14, bulletFontName=uni_font, bulletFontSize=10.5))
                continue
            if lines[i].strip():
                buf.append(lines[i])
            else:
                flush()
            i += 1
        flush()
        return flow

    def md_to_flowables(text: str):
        flow, pos = [], 0
        for m in code_fence.finditer(text):
            flow += paragraphs_and_lists(text[pos:m.start()])
            flow.append(Preformatted(m.group(1).replace("\t", "    "), style=code_style))
            pos = m.end()
        return flow + paragraphs_and_lists(text[pos:])

    flow = [Paragraph("Law Jublie — Chat Export", title_style)]
    for msg in messages:
        flow.append(Paragraph(xml_escape(str(msg.get("role", "assistant")).capitalize()), role_style))
        flow += md_to_flowables(str(msg.get("content", "")))
        flow.append(Spacer(1, 6))

    pdf_buf = BytesIO()
    SimpleDocTemplate(pdf_buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm,
                      topMargin=16 * mm, bottomMargin=16 * mm).build(flow)
    st.download_button("📄 Download as PDF", data=pdf_buf.getvalue(),
                       file_name="law_chat.pdf", mime="application/pdf")


# ------------------------------
# Fine lookup
# ------------------------------
def search_fine_amount_india(user_query: str) -> dict | None:
    """Asks the LLM for the specific Indian fine/penalty for an offence. Returns a dict or None."""
    sys_prompt = (
        "You are a legal assistant specialized in Indian law. "
        "Given a user question about an offence, identify the specific law/section and state the current fine/penalty "
        "amount in India with currency and a short source note. "
        "Return ONLY valid JSON with keys: amount, currency, section, act, source. "
        "If unknown, return an empty JSON object {}."
    )
    try:
        data = extract_json(chat([
            {"role": "system", "content": sys_prompt},
            {"role": "user", "content": f"India {user_query}"},
        ], temperature=0.1, max_tokens=1200))
    except Exception:
        return None
    if not data or not data.get("amount"):
        return None
    return {k: data.get(k) for k in ("amount", "currency", "section", "act", "source")}


def _format_fine(fine: dict) -> str:
    currency = f" ({fine['currency']})" if fine.get("currency") else ""
    lines = ["\n\n💸 **Fine (India):**", f"- **Amount:** {fine['amount']}{currency}"]
    lines += [f"- **{label}:** {fine[key]}" for key, label in
              (("section", "Section"), ("act", "Act"), ("source", "Source")) if fine.get(key)]
    return "\n".join(lines)


CASE_HELP_MD = """
### 🔍 Step 1 — Open eCourts Website
Use the official government portal (link below), or search "eCourts India".

---

### ✅ Step 2 — Choose Search Method

• CNR Number (Best & fastest)
• Case Number
• Party Name
• Advocate Name
• FIR Number

---

### 🏛 Step 3 — Select Location

State → District → Court Complex

⚠️ Wrong district = no result

---

### 📄 Step 4 — Case Number Format

Case Type — CC / CRL / WP / OS
Case Number — numeric
Year — filing year

Example: CC / 123 / 2022

---

### 🚓 FIR Search Needs

• FIR number
• Police station
• Year

---

### 🔐 Step 5 — Enter Captcha

Type exactly → Click Search

---

### 📊 Results Show

✅ Case status
✅ Hearing dates
✅ Orders
✅ Judgments
✅ Judge details

---

### 💡 Tips

✔ Use CNR when available
✔ Check case year carefully
✔ Try partial names
✔ Avoid spelling mistakes

NOTE: Use Com-Jublie for help understanding legal terms, not for case status.
"""


def _add_assistant(msg: str):
    st.session_state.law_messages.append({"role": "assistant", "content": msg})
    st.chat_message("assistant").markdown(msg)


# ------------------------------
# Main Chatbot UI
# ------------------------------
def show_law_chatbot():
    st.subheader("⚖️ Law Jublie — Constitutional & IPC Legal Assistant")

    st.session_state.setdefault("law_messages", [])
    st.session_state.setdefault("show_complaint_form", False)
    st.session_state.setdefault("show_case_help", False)
    st.session_state.setdefault("processed_pdf", None)

    with st.expander("💾 Export Chat History"):
        col1, col2 = st.columns(2)
        with col1:
            if st.button("📄 Export as PDF"):
                export_chat("pdf")
        with col2:
            if st.button("📘 Export as DOCX"):
                export_chat("docx")

    st.divider()

    for msg in st.session_state.law_messages:
        with st.chat_message(msg["role"]):
            st.markdown(msg["content"])

    # PDF upload — summarize each file once, not on every rerun
    pdf_file = st.file_uploader("📂 Upload a law-related PDF", type=["pdf"])
    if pdf_file and st.session_state.processed_pdf != pdf_file.file_id:
        st.session_state.processed_pdf = pdf_file.file_id
        with st.spinner("🔍 Extracting and analyzing your document..."):
            extracted_text = read_pdf(pdf_file)
            if not extracted_text:
                st.warning("⚠️ Could not extract any readable text from the PDF.")
            elif not is_law_related(extracted_text):
                _add_assistant("⚠️ This PDF doesn't seem related to Indian law.")
            else:
                _add_assistant(generate_summary(extracted_text))

    st.divider()

    col1, col2, col3 = st.columns([7, 2, 2])
    with col1:
        law_query = st.chat_input("💬 Ask a law-related question...", key="law_input")
    with col2:
        if st.button("📢 File Complaint"):
            st.session_state.show_complaint_form = not st.session_state.show_complaint_form
    with col3:
        if st.button("📂 Check Case History"):
            st.session_state.show_case_help = True

    if st.session_state.show_case_help:
        st.markdown("---")
        st.markdown("## 📘 How to Check Case Details — eCourts Guide")
        with st.container(border=True):
            st.markdown(CASE_HELP_MD)
            st.link_button("🔍 Open eCourts Website", "https://services.ecourts.gov.in/ecourtindia_v6/")
            if st.button("❌ Close Guide"):
                st.session_state.show_case_help = False
                st.rerun()
        st.markdown("---")

    if st.session_state.show_complaint_form:
        st.subheader("📝 Submit a Complaint")
        with st.form("complaint_form"):
            title = st.text_input("Complaint Title")
            description = st.text_area("Complaint Description")
            if st.form_submit_button("Submit"):
                if title.strip() and description.strip():
                    with open(COMPLAINTS_FILE, "a", encoding="utf-8") as f:
                        f.write(f"Title: {title}\nDescription: {description}\n\n")
                    st.success("✅ Complaint submitted successfully.")
                    st.session_state.show_complaint_form = False
                else:
                    st.error("⚠️ Both title and description are required.")

    if not law_query:
        return

    st.session_state.law_messages.append({"role": "user", "content": law_query})
    with st.chat_message("user"):
        st.markdown(law_query)

    if not is_law_related(law_query):
        _add_assistant("❗ This chatbot only handles Indian law queries (Constitution, IPC/BNS, and other Acts).")
        return

    with st.spinner("⚖️ Jublie is analyzing your query..."):
        system_prompt = (
            "You are 'Jublie', a senior legal advisor specialized in Indian law. "
            "Give concise, factual answers grounded in the Constitution of India and Indian statutes. "
            "For criminal law, cite the Bharatiya Nyaya Sanhita (BNS) / BNSS / BSA, which replaced the IPC / CrPC / "
            "Evidence Act from 1 July 2024, and mention the old IPC/CrPC section where helpful. "
            "Never invent section numbers; if unsure, say so. Avoid opinions or assumptions."
        )
        try:
            reply = chat([
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": f"India {law_query}"},
            ], model=CHAT_MODEL, temperature=0.3, max_tokens=1024)
        except Exception as e:
            st.error(str(e))
            return

        # Verified text for any Constitution Article the user mentioned
        lookup = coi_lookup()
        for ref in article_refs(law_query)[:3]:
            if ref in lookup:
                reply += f"\n\n📘 **Verified Legal Reference — Article {ref.upper()}:**\n{lookup[ref][:1500]}"

        if _FINE_TERMS.search(law_query):
            fine_info = search_fine_amount_india(law_query)
            if fine_info:
                reply += _format_fine(fine_info)

        _add_assistant(reply)
